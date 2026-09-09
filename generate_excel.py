import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Student Data Structure
data = [
    # Grade 10 - Philippians
    {
        "grade_section": "Grade 10 – Philippians",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "Jeanne S. Mirayo, LPT",
        "students": [
            ("Bangcola, Aiman C.", "", "Allowed", "Blue"),
            ("De Los Santos, Rodrigo III O.", "", "Allowed", "Blue"),
            ("Laroza, Rex Julian T.", "", "Allowed", "Blue"),
            ("Nacilla, Miguel Kriztoffe S.", "", "Allowed", "Blue"),
            ("Garcia, Alexander Ross A.", "", "Allowed", "Blue"),
            ("Salangsang, Alby Kristian B.", "", "Allowed", "Blue"),
            ("Malasador, John Henry S.", "", "Not allowed", "Red"),
            ("Undangan, John David T.", "", "Allowed", "Blue"),
            ("Coronilia, Paul Jacob G.", "", "Allowed at the waiting area", "White"),
            ("Signey Jr., Ronald Allan C.", "", "Allowed", "Blue"),
            ("Palec, Mecan James A.", "", "Allowed", "Blue"),
            ("Besin, Don Gabriel L.", "", "Allowed", "Blue"),
            ("De Leon, Edara Charis A.", "", "Allowed", "Blue"),
            ("Lao, Sabelle Monica L.", "", "Allowed", "Blue"),
            ("Salangsang, Lyrica Nehtania V.", "", "Allowed", "Blue"),
            ("Sandagon, Synne Klarizz V.", "", "Allowed", "Blue"),
            ("Tan, Mischa Arabella S.", "", "Allowed", "Blue"),
            ("Cabanlit, Anika Reese P.", "", "Allowed", "Blue"),
            ("Ballena, Carmina Alessandra J.", "", "Allowed", "Blue"),
            ("Donquines, Sophia Jeasie D.", "", "Allowed", "Blue"),
            ("Deita, Promise Q.", "", "Not allowed", "Red"),
            ("Movilla, Julia Elliana L.", "", "Allowed", "Blue"),
            ("Tan, Ticiaviel D.", "", "Allowed", "Blue"),
            ("Gonzales, Eunice C.", "", "Allowed", "Blue"),
            ("Quirino, Cayleigh Yj J.", "", "Allowed at the waiting area", "White"),
            ("Forones, Christine Mae B.", "", "Allowed", "Blue"),
            ("Santos, Ella Grace L.", "", "Allowed", "Blue"),
            ("Coronado, Victoria Marie B.", "", "Allowed", "Blue"),
        ]
    },
    # Grade 3 - Andrew
    {
        "grade_section": "Grade 3 – Andrew",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "",
        "students": [
            ("Aniversario, Isaiah Miguel N.", "Boy", "Not allowed", "Red"),
            ("Chio, Kyler Matthew", "Boy", "Allowed at the waiting area", "White"),
            ("Delfin, Christensen B.", "Boy", "Unmarked", "-"),
            ("Nierra, Matth Levi G.", "Boy", "Not allowed", "Red"),
            ("Olamit, Oni Marko C.", "Boy", "Not allowed", "Red"),
            ("Rudwick, Griffin Ernest A.", "Boy", "Allowed", "Blue"),
            ("Sorilla, Phil Sebastian R.", "Boy", "Allowed", "Blue"),
            ("Teves, Marco Alfonso S.", "Boy", "Allowed", "Blue"),
            ("Umadhay, Kino Brenner S.", "Boy", "Not allowed", "Red"),
            ("Yap, Victor James C. III", "Boy", "Not allowed", "Red"),
            ("Yu, Tyron Raphael Y.", "Boy", "Allowed", "Blue"),
            ("Agustino, Janina Camila N.", "Girl", "Not allowed", "Red"),
            ("Alabado, Margaux Mirana D.", "Girl", "Not allowed", "Red"),
            ("Bacalso, Helyan Vitelli Cecilia G.", "Girl", "Not allowed", "Red"),
            ("Chua, Nancy K.", "Girl", "Unmarked", "-"),
            ("Crave, Rimona Jessica L.", "Girl", "Unmarked", "-"),
            ("Dawang, Tammie Dominique I.", "Girl", "Allowed", "Blue"),
            ("Jamolin, Florein Ahrea L.", "Girl", "Allowed", "Blue"),
            ("Leggett, Imogen Angel Joy A.", "Girl", "Allowed", "Blue"),
            ("Lorenzo, Kristina Mithi N.", "Girl", "Allowed", "Blue"),
            ("Mirabueno, Bella Caitlyn L.", "Girl", "Allowed", "Blue"),
            ("Olarte, Summer Amira K.", "Girl", "Not allowed", "Red"),
            ("Sabado, Aila Sophia", "Girl", "Allowed at the waiting area", "White"),
            ("Sabatin, Jillianne Martine C.", "Girl", "Allowed", "Blue"),
            ("Salvilla, Alia Khadija S.", "Girl", "Not allowed", "Red"),
            ("Tan, Audrey Samantha I.", "Girl", "Allowed", "Blue"),
            ("Yuson, Gabrielle Dawn N.", "Girl", "Allowed", "Blue"),
        ]
    },
    # Grade 3 - Peter
    {
        "grade_section": "Grade 3 – Peter",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "",
        "students": [
            ("De Vera, Linus Ezekiel R.", "Boy", "Not allowed", "Red"),
            ("Franco, Kyle Nathaniel O.", "Boy", "Not allowed", "Red"),
            ("Frolov, Vasili Victorovich", "Boy", "Allowed at the waiting area", "White"),
            ("Lopez, Zhab B.", "Boy", "Not allowed", "Red"),
            ("Posadas, Dick Zion L.", "Boy", "Not allowed", "Red"),
            ("Regalado, Liandre Phylix Q.", "Boy", "Allowed at the waiting area", "White"),
            ("Teves, Liam Gil T.", "Boy", "Not allowed", "Red"),
            ("Tidula, Homer Al-Bryan L.", "Boy", "Not allowed", "Red"),
            ("Torres, Alexius Adrielle Z.", "Boy", "Allowed", "Blue"),
            ("Alabado, Meghan Kelsey D.", "Girl", "Not allowed", "Red"),
            ("Aranas, Eesha Yuan E.", "Girl", "Allowed", "Blue"),
            ("Belleguez, Gzrel Kate L.", "Girl", "Not allowed", "Red"),
            ("Cataluña, Coraline Celestin D.", "Girl", "Allowed", "Blue"),
            ("Cawa, Jeanna Athena C.", "Girl", "Allowed at the waiting area", "White"),
            ("Cho, Yoona E.", "Girl", "Allowed", "Blue"),
            ("Dela Cruz, Sheena Faye H.", "Girl", "Not allowed", "Red"),
            ("Garcia, Mari Ysobelle P.", "Girl", "Not allowed", "Red"),
            ("Gillett, Abby Jane M.", "Girl", "Not allowed", "Red"),
            ("Intong, Nicole Bryce M.", "Girl", "Allowed", "Blue"),
            ("Lumawag, Avana Abrienne D.", "Girl", "Not allowed", "Red"),
            ("Lumaya, Xybelle Blythe B.", "Girl", "Not allowed", "Red"),
            ("Matias, H Joey Feeve A.", "Girl", "Not allowed", "Red"),
            ("Ramos, Adrielle Hope A.", "Girl", "Allowed", "Blue"),
            ("Solde, Yasha Symione H.", "Girl", "Allowed at the waiting area", "White"),
            ("Sun, Nina Olivia C.", "Girl", "Allowed", "Blue"),
            ("Tababa, Liana Pia C.", "Girl", "Not allowed", "Red"),
            ("Torres, Avanna Adreanna Z.", "Girl", "Allowed", "Blue"),
            ("Uy, Xia Michaela M.", "Girl", "Not allowed", "Red"),
        ]
    },
    # Grade 4 - James
    {
        "grade_section": "Grade 4 – James",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "Ivy Q. Maquinay, LPT",
        "students": [
            ("Besin, Don Samuel R.", "Boy", "Allowed", "Blue"),
            ("Dela Cruz, Anton Miguel R.", "Boy", "Not allowed", "Red"),
            ("Duropan, Mharlie Seanan U.", "Boy", "Allowed", "Blue"),
            ("Guo, James Cian L.", "Boy", "Allowed", "Blue"),
            ("Isaloc, Rheymer John A.", "Boy", "Not allowed", "Red"),
            ("Maniago, Erj David D.", "Boy", "Allowed", "Blue"),
            ("Mindalanao, Muhammad Azhar L.", "Boy", "Allowed", "Blue"),
            ("Palacio, Jedrile Ethan E.", "Boy", "Allowed at the waiting area", "White"),
            ("Ramos, David Paul D.", "Boy", "Not allowed", "Red"),
            ("Ramos, Matthew Ryan A.", "Boy", "Allowed", "Blue"),
            ("Roque, Hezekiah Jerahmeel S.", "Boy", "Allowed", "Blue"),
            ("Venturina, Caleb Gabriel S.", "Boy", "Allowed", "Blue"),
            ("Villaverde, Joaquin Andres A.", "Boy", "Allowed", "Blue"),
            ("Yu, Stephan Amrit M.", "Boy", "Not allowed", "Red"),
            ("Cahayag, Yana Zafina S.", "Girl", "Allowed", "Blue"),
            ("Estabillo, Hezekel B.", "Girl", "Allowed at the waiting area", "White"),
            ("Fabila, Giana Edrisse C.", "Girl", "Not allowed", "Red"),
            ("Gravidez, Jilliane Louise", "Girl", "Allowed", "Blue"),
            ("Hamac, Margaux Elyce C.", "Girl", "Allowed", "Blue"),
            ("Herrera, Dee Audrey Z.", "Girl", "Allowed", "Blue"),
            ("Lacson, Sumin P.", "Girl", "Allowed", "Blue"),
            ("Leyson, Abrianna Eleisha M.", "Girl", "Not allowed", "Red"),
            ("Linzag, Bai Hasanat K.", "Girl", "Not allowed", "Red"),
            ("Ongayo, Alexandria A.", "Girl", "Not allowed", "Red"),
            ("Palma, Ronna Isabelle J.", "Girl", "Not allowed", "Red"),
            ("Pigar, Katarina Brielle T.", "Girl", "Not allowed", "Red"),
            ("Rogero, Maria Cresia P.", "Girl", "Not allowed", "Red"),
            ("Yupengco, Anne Kaitlin G.", "Girl", "Allowed at the waiting area", "White"),
        ]
    },
    # Grade 5 - Aaron
    {
        "grade_section": "Grade 5 – Aaron",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "Georiza Mae R. Anguay, LPT",
        "students": [
            ("Ablog, Deandrei James D.", "Boy", "Allowed", "Blue"),
            ("Ablog, Zeandrew James D.", "Boy", "Allowed", "Blue"),
            ("Azucena, Xalm Noah G.", "Boy", "Not allowed", "Red"),
            ("Crave, Renzo Jess L.", "Boy", "Allowed", "Blue"),
            ("Faeldonia, John Zohan M.", "Boy", "Allowed", "Blue"),
            ("Francisco, Gabrielle Andrew D.", "Boy", "Allowed", "Blue"),
            ("Lara, Mikieljed W.", "Boy", "Allowed", "Blue"),
            ("Linzmaier, Ethan Andrew C.", "Boy", "Not allowed", "Red"),
            ("Mancera, Luis Zandrew R.", "Boy", "Allowed at the waiting area", "White"),
            ("Martinez, Michael Laurent N.", "Boy", "Allowed at the waiting area", "White"),
            ("Mendiola, Heuer A.", "Boy", "Allowed", "Blue"),
            ("Snalam, Charles R.", "Boy", "Not allowed", "Red"),
            ("Tatualla, Ziev Undrei C.", "Boy", "Not allowed", "Red"),
            ("Torres, Tristan Cyan M.", "Boy", "Not allowed", "Red"),
            ("Valencia, Deandre Leighton T.", "Boy", "Allowed at the waiting area", "White"),
            ("Aniversario, Giuliana Beatrice N.", "Girl", "Allowed at the waiting area", "White"),
            ("Bacasdoon, Celeste Reign D.", "Girl", "Allowed", "Blue"),
            ("Chua, Yuwen P.", "Girl", "Allowed", "Blue"),
            ("Dinopol, Ma. Olivia Soleil P.", "Girl", "Allowed", "Blue"),
            ("Egasan, Zephaniah M.", "Girl", "Not allowed", "Red"),
            ("Galanza, Shaelee Agatha S.", "Girl", "Not allowed", "Red"),
            ("Gillet, Kristin Leah M.", "Girl", "Not allowed", "Red"),
            ("Hung, Penelope Eve M.", "Girl", "Allowed", "Blue"),
            ("Khow, Ceion F.", "Girl", "Allowed", "Blue"),
            ("Malveaux, Ashira A.", "Girl", "Allowed", "Blue"),
            ("Ng, Wai Ching", "Girl", "Allowed at the waiting area", "White"),
            ("Pacquiao, Ruella Ysabel R.", "Girl", "Allowed", "Blue"),
            ("Santos, Elexa Grace L.", "Girl", "Allowed", "Blue"),
            ("Tapel, Shian Riley A.", "Girl", "Allowed", "Blue"),
            ("Yap, Sky Gabrielle R.", "Girl", "Allowed", "Blue"),
        ]
    },
    # Grade 5 - Moses
    {
        "grade_section": "Grade 5 – Moses",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "Princess Claire A. Bacolod, LPT",
        "students": [
            ("Aparente, Jedidiah Elisus D.", "Boy", "Allowed", "Blue"),
            ("Aparente, Joziah Elisus D.", "Boy", "Allowed", "Blue"),
            ("Barranco, Jacob Clyde B.", "Boy", "Allowed at the waiting area", "White"),
            ("Bation, Lyx Xcyljy D.", "Boy", "Not allowed", "Red"),
            ("Bolongon, Francis Zian J.", "Boy", "Not allowed", "Red"),
            ("Cimini, Searle Joab M.", "Boy", "Allowed at the waiting area", "White"),
            ("Gimeno, Kenzo P.", "Boy", "Allowed", "Blue"),
            ("Guo, Johann Caden L.", "Boy", "Allowed", "Blue"),
            ("Perales, Lucas Oliver D.", "Boy", "Allowed", "Blue"),
            ("Rojo, ELijah Dvans R.", "Boy", "Not allowed", "Red"),
            ("Sandagon, Scott V.", "Boy", "Allowed", "Blue"),
            ("Sevilla, Joaquin Kristoff C.", "Boy", "Allowed", "Blue"),
            ("Villanueva, Julien Cael D.", "Boy", "Allowed", "Blue"),
            ("Yuson, Genesis Den N.", "Boy", "Allowed", "Blue"),
            ("Agustino, Jeiah Carmela N.", "Girl", "Not allowed", "Red"),
            ("Anas, Serena Gabrielle B.", "Girl", "Not allowed", "Red"),
            ("Aranas, Hyacinth yuan E.", "Girl", "Allowed", "Blue"),
            ("Bacasdoon, Aden Claire F.", "Girl", "Allowed", "Blue"),
            ("Blancada, Felicity Rose P.", "Girl", "Not allowed", "Red"),
            ("Cornel, Ashlee Elise A.", "Girl", "Allowed", "Blue"),
            ("Cortez, Mary Nathalia G.", "Girl", "Not allowed", "Red"),
            ("Dizon, Erin Sofia J.", "Girl", "Allowed", "Blue"),
            ("Esma, Morghana Grace O.", "Girl", "Not allowed", "Red"),
            ("Flores, Joleen F.", "Girl", "Allowed", "Blue"),
            ("Lagudas, Marguerette D.", "Girl", "Not allowed", "Red"),
            ("Laroza, Alexandria T.", "Girl", "Allowed at the waiting area", "White"),
            ("Lumaque, Riae Valkyrie Y.", "Girl", "Allowed", "Blue"),
            ("Quianzon, Elaiza Andrea A.", "Girl", "Allowed", "Blue"),
            ("Roldan, Kezhea Dawn", "Girl", "Allowed at the waiting area", "White"),
            ("Ybañez, Ziana Monique", "Girl", "Allowed", "Blue"),
        ]
    },
    # Grade 7 - Galatians
    {
        "grade_section": "Grade 7 – Galatians",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "Irish T. Dayondon, LPT",
        "students": [
            ("Borreros, Hosea Wayne I.", "Boy", "Allowed", "Blue"),
            ("Cape, Mateo Iñigo Q.", "Boy", "Allowed", "Blue"),
            ("Erbina, Augustus Miguel D.", "Boy", "Allowed", "Blue"),
            ("Estenor, Eman Jae L.", "Boy", "Allowed", "Blue"),
            ("Felonia, Leonardo F. Jr", "Boy", "Unmarked", "-"),
            ("Go, Jenson U.", "Boy", "Unmarked", "-"),
            ("Movilla, Jake Ernest L.", "Boy", "Allowed", "Blue"),
            ("Pia, Azarel Huey B.", "Boy", "Allowed", "Blue"),
            ("Selvido, Lucas Izak S.", "Boy", "Allowed", "Blue"),
            ("Yupengco, Kerwyn G.", "Boy", "Allowed", "Blue"),
            ("Anoche, Shane Ashlei F.", "Girl", "Allowed", "Blue"),
            ("Bendero, Amanda Feliz V.", "Girl", "Allowed", "Blue"),
            ("Buisan, Elyannah C.", "Girl", "Allowed", "Blue"),
            ("Dela Cruz, Angela Gabrielle R.", "Girl", "Not allowed", "Red"),
            ("Diono, Sofiya Allaira A.", "Girl", "Allowed", "Blue"),
            ("Hong, Catherine A.", "Girl", "Allowed", "Blue"),
            ("Lu, Coleen Gabrianne K.", "Girl", "Allowed", "Blue"),
            ("Morin, Denisse T.", "Girl", "Not allowed", "Red"),
            ("Pacquiao, Eira Robelle R.", "Girl", "Allowed", "Blue"),
            ("Ruiz, Pia Dominique T.", "Girl", "Allowed at the waiting area", "White"),
            ("Tapel, Alexis Madison A.", "Girl", "Allowed", "Blue"),
            ("Tingson, Mary Ayeesha C.", "Girl", "Allowed", "Blue"),
            ("Undangan, Bernice Kaye T.", "Girl", "Allowed at the waiting area", "White"),
        ]
    },
    # Grade 8 - Leviticus
    {
        "grade_section": "Grade 8 – Leviticus",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "Merlaine Gay G. Jacaba, LPT",
        "students": [
            ("Andaya, Kiefer D.", "Boy", "Allowed", "Blue"),
            ("Awayid, Ameen D.", "Boy", "Allowed", "Blue"),
            ("Cahayag, Andrei Benedict S.", "Boy", "Allowed", "Blue"),
            ("Calong, Johan Jeiel F.", "Boy", "Allowed", "Blue"),
            ("Chatto, Cliff Martin E.", "Boy", "Not allowed", "Red"),
            ("Chio, Johann Theodore Emmanuel A.", "Boy", "Allowed", "Blue"),
            ("Coronado, Ethan B.", "Boy", "Allowed", "Blue"),
            ("Dumalag, David Miguel S.", "Boy", "Allowed at the waiting area", "White"),
            ("Fernandez, Seth Aiden T.", "Boy", "Allowed", "Blue"),
            ("Leonor, Russel John E.", "Boy", "Allowed", "Blue"),
            ("Nadua, Rare R.", "Boy", "Allowed", "Blue"),
            ("Salangsang, Lyric Lyiam V.", "Boy", "Allowed at the waiting area", "White"),
            ("Sampilo, John Sigfrid M.", "Boy", "Not allowed", "Red"),
            ("Suelto, Lexan Dave A.", "Boy", "Allowed", "Blue"),
            ("Te, Max Benedict H.", "Boy", "Allowed", "Blue"),
            ("Tuazon, Andrew Josh S.", "Boy", "Allowed", "Blue"),
            ("Yu, Sebastian Anick M.", "Boy", "Unmarked", "-"),
            ("Abañez, Breena Marthina M.", "Girl", "Allowed", "Blue"),
            ("Apsay, Elif O.", "Girl", "Allowed", "Blue"),
            ("Arias, Ella Valerie C.", "Girl", "Allowed at the waiting area", "White"),
            ("Cabading, Jabeth Neb R.", "Girl", "Allowed", "Blue"),
            ("Camarillo, Ellish Mari R.", "Girl", "Allowed", "Blue"),
            ("Congson, Aimelyn P.", "Girl", "Allowed", "Blue"),
            ("De Los Santos, Rashina O.", "Girl", "Allowed", "Blue"),
            ("Tidula, Rhohann Marie L.", "Girl", "Allowed", "Blue"),
            ("Vicario, Sofia Khloe A.", "Girl", "Allowed", "Blue"),
        ]
    },
    # Grade 8 - Numbers
    {
        "grade_section": "Grade 8 – Numbers",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "",
        "students": [
            ("Anoche, Shone Andrei F.", "Boy", "Allowed", "Blue"),
            ("Bernales, Yosef Caleb C.", "Boy", "Allowed", "Blue"),
            ("Chua, Paul Kian P.", "Boy", "Allowed", "Blue"),
            ("Egasan, Kaze Brione E.", "Boy", "Allowed at the waiting area", "White"),
            ("Magsipoc, Pierre Aldrei G.", "Boy", "Allowed", "Blue"),
            ("Maturan, Euan M.", "Boy", "Allowed", "Blue"),
            ("Pabillo, Neil Anthony C.", "Boy", "Allowed", "Blue"),
            ("Pacete, Francis Joaquin C.", "Boy", "Allowed", "Blue"),
            ("Palacio, Jed Edrian E.", "Boy", "Allowed", "Blue"),
            ("Qiu, Yuanxin", "Boy", "Allowed", "Blue"),
            ("Sales, Mikael Clyde", "Boy", "Allowed", "Blue"),
            ("Samonte, Jethro Franco", "Boy", "Allowed at the waiting area", "White"),
            ("Siason, Jearn Adge D.", "Boy", "Allowed", "Blue"),
            ("Suertigosa, Aldwyne Mar C.", "Boy", "Allowed", "Blue"),
            ("Tabar, Luis Iñigo D.", "Boy", "Allowed", "Blue"),
            ("Ang, Elianah Faith B.", "Girl", "Allowed at the waiting area", "White"),
            ("Aureo, Julia P.", "Girl", "Allowed", "Blue"),
            ("Chu, Shanel Venice", "Girl", "Allowed", "Blue"),
            ("Dela Cruz, Alyana Isabel C.", "Girl", "Allowed", "Blue"),
            ("Escobar, Daenerys Arya A.", "Girl", "Allowed", "Blue"),
            ("Gerada, Azrielle Nesia J.", "Girl", "Allowed", "Blue"),
            ("Jamolin, Hesthea Shaey O.", "Girl", "Allowed", "Blue"),
            ("Lopez, Aehra Beth B.", "Girl", "Allowed", "Blue"),
            ("Macalipay, Jellian Clate I", "Girl", "Allowed at the waiting area", "White"),
            ("Octavio, Karylle Grace Q.", "Girl", "Allowed", "Blue"),
            ("Ong, Kalia Annika C.", "Girl", "Allowed", "Blue"),
            ("Ramos, Margareth Ashley P.", "Girl", "Not allowed", "Red"),
            ("Sunglao, Sapphia Monique O.", "Girl", "Allowed", "Blue"),
        ]
    },
    # Grade 9 - Proverbs
    {
        "grade_section": "Grade 9 – Proverbs",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "George Kim A. Barcelona, LPT",
        "students": [
            ("Alegario, Geoff Euhanz S.", "Boy", "Allowed", "Blue"),
            ("Asuncion, Jon Rino Miguel F.", "Boy", "Allowed", "Blue"),
            ("Banaynal, Andrie Stefan M.", "Boy", "Allowed", "Blue"),
            ("Borres, Carl Joseph P.", "Boy", "Allowed", "Blue"),
            ("Cahayag, Sean Cedric S.", "Boy", "Allowed", "Blue"),
            ("DE GUZMAN, SEBASTIEN HUGO A.", "Boy", "Allowed", "Blue"),
            ("Linzag, Nur Amin K.", "Boy", "Allowed", "Blue"),
            ("NACIONAL, JACOB A.", "Boy", "Allowed", "Blue"),
            ("Oleta, Jheron L.", "Boy", "Allowed", "Blue"),
            ("REHFELDT, FRANK D.", "Boy", "Allowed at the waiting area", "White"),
            ("Relacion, Dennie Louie L.", "Boy", "Allowed", "Blue"),
            ("Roda, Red Lawrence J.", "Boy", "Allowed", "Blue"),
            ("RUGA, JACOB RHANIEL G.", "Boy", "Allowed", "Blue"),
            ("Tapaya, Vander Bry B.", "Boy", "Allowed", "Blue"),
            ("Torcelino, Ethan Leon Jaime S.", "Boy", "Allowed", "Blue"),
            ("Villanueva, Liam Piotrek L.", "Boy", "Allowed", "Blue"),
            ("Alabado, Macy Mariana D.", "Girl", "Allowed", "Blue"),
            ("Baconguis, Samantha Dane D.", "Girl", "Allowed", "Blue"),
            ("Baraquel, Sacha Victoria Z.", "Girl", "Allowed", "Blue"),
            ("Dinopol, Evanna Denise", "Girl", "Allowed", "Blue"),
            ("Lagudas, Florence Kaileen N.", "Girl", "Allowed", "Blue"),
            ("LASPIÑAS, KHREA SHENNE P.", "Girl", "Allowed", "Blue"),
            ("Linzmaier, Desiree Nicole C.", "Girl", "Allowed", "Blue"),
            ("Ong, Adreana Skye O.", "Girl", "Allowed", "Blue"),
            ("PE, SUYENN R.", "Girl", "Allowed", "Blue"),
            ("SANGGO, SOPHIA CARLEE Q.", "Girl", "Not allowed", "Red"),
            ("Tan, Amia Noelle B.", "Girl", "Allowed", "Blue"),
            ("Tan, Colleen Erika Mikaella U.", "Girl", "Allowed", "Blue"),
            ("WONG, LEILA MICHELLE N.", "Girl", "Allowed", "Blue"),
            ("Yu, Claudette C.", "Girl", "Allowed", "Blue"),
        ]
    },
    # Grade 9 - Psalms
    {
        "grade_section": "Grade 9 – Psalms",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "Kriszel Joy C. Ramos, LPT",
        "students": [
            ("Caga, Mico", "Boy", "Allowed", "Blue"),
            ("Dy, Neil John S.", "Boy", "Allowed", "Blue"),
            ("Escleto, James Roger T.", "Boy", "Allowed at the waiting area", "White"),
            ("Hong, Yingjian S.", "Boy", "Allowed", "Blue"),
            ("Layog, Rhian Gercy L.", "Boy", "Allowed", "Blue"),
            ("Macalipay, Zian Jilary I.", "Boy", "Allowed at the waiting area", "White"),
            ("Mathong, Ace A.", "Boy", "Allowed", "Blue"),
            ("Nix, John Paulo A.", "Boy", "Allowed", "Blue"),
            ("Ong, Jefferson Patrick V.", "Boy", "Allowed", "Blue"),
            ("Ortiz, Conrad P.", "Boy", "Allowed", "Blue"),
            ("Pinalgan, Prince Peter Priel C.", "Boy", "Allowed", "Blue"),
            ("Provido, Nathan Xander", "Boy", "Allowed", "Blue"),
            ("Tandoy, Kyt Daniel B.", "Boy", "Allowed", "Blue"),
            ("Bandrang, Ghashia Daneen Lucman S.", "Girl", "Allowed", "Blue"),
            ("Cabanlit, Seychelles Rucci", "Girl", "Allowed", "Blue"),
            ("Cerado, Juliana Mariz Angelica T.", "Girl", "Allowed", "Blue"),
            ("Chua, Yiwen P.", "Girl", "Allowed", "Blue"),
            ("Clarin, Christine Jireh A.", "Girl", "Allowed", "Blue"),
            ("Festin, Cha Dixie M.", "Girl", "Allowed", "Blue"),
            ("Hong, Feona H.", "Girl", "Allowed", "Blue"),
            ("Leron, Prezeah Darylie A.", "Girl", "Allowed", "Blue"),
            ("Lozano, Jaisha", "Girl", "Allowed", "Blue"),
            ("Pelicano, Michailla Angela L.", "Girl", "Allowed", "Blue"),
            ("Pontino, Amanda Celina S.", "Girl", "Allowed", "Blue"),
            ("Salcedo, Acee Jane P.", "Girl", "Allowed", "Blue"),
            ("Sarno, Amber A.", "Girl", "Allowed", "Blue"),
            ("Tan, Jessrich Ambria M.", "Girl", "Allowed", "Blue"),
            ("Uy, Beatrice R.", "Girl", "Allowed", "Blue"),
            ("Velasquez, Riona Aneesa L.", "Girl", "Allowed", "Blue"),
            ("Young, Cassandra Ashley B.", "Girl", "Allowed", "Blue"),
        ]
    },
    # Grade 10 - Ephesians
    {
        "grade_section": "Grade 10 – Ephesians",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "",
        "students": [
            ("Buisan, Ethan C.", "Boy", "Allowed", "Blue"),
            ("Cariazo, Astaro Kerio M.", "Boy", "Allowed", "Blue"),
            ("Creedon, William Scott B.", "Boy", "Allowed", "Blue"),
            ("Diono, Thomas Allai A.", "Boy", "Allowed", "Blue"),
            ("Fabila, Rogie C.", "Boy", "Allowed", "Blue"),
            ("Lobaton, Zeoh James V.", "Boy", "Allowed", "Blue"),
            ("Mercado, Rizly Carl F.", "Boy", "Allowed at the waiting area", "White"),
            ("Ong, Aiken P.", "Boy", "Allowed", "Blue"),
            ("Paches, Jose Ricardo Antonio T.", "Boy", "Allowed", "Blue"),
            ("Panlaque, Thevron Vicam F.", "Boy", "Allowed", "Blue"),
            ("Sanchez, Sean Philippe C.", "Boy", "Allowed", "Blue"),
            ("She, Jin Nuo P.", "Boy", "Allowed", "Blue"),
            ("Subaldo, Adam Lewis F.", "Boy", "Allowed", "Blue"),
            ("Torcelino, Julian Matteo S.", "Boy", "Allowed", "Blue"),
            ("Villanueva, Regie B.", "Boy", "Allowed at the waiting area", "White"),
            ("Abañez, Brianna Matthea M.", "Girl", "Allowed", "Blue"),
            ("Baltazar, Katara Atasha", "Girl", "Allowed", "Blue"),
            ("Bercero, Adelaine Faye C.", "Girl", "Not allowed", "Red"),
            ("Clavite, Kerstin Claude O.", "Girl", "Allowed", "Blue"),
            ("Duarte, Stephanie Celeste L.", "Girl", "Allowed", "Blue"),
            ("Escobar, Dimee Miguella A.", "Girl", "Allowed", "Blue"),
            ("Jones, Halle T.", "Girl", "Allowed", "Blue"),
            ("Lu, Chelsea Gwen K.", "Girl", "Allowed", "Blue"),
            ("Quesada, Ysabella Sophia B.", "Girl", "Allowed", "Blue"),
            ("Salvador, Cheska Nicole V.", "Girl", "Allowed", "Blue"),
            ("Tan, Casey Ysabelle T.", "Girl", "Allowed", "Blue"),
        ]
    },
    # Grade 12 - Isaac and Jacob
    {
        "grade_section": "Grade 12 – Isaac and Jacob",
        "school_year": "S.Y. 2025 – 2026",
        "adviser": "Queenylaine B. Prieto, LPT",
        "students": [
            ("Albinda, Xeb Emmanuel A.", "Boy", "Allowed", "Blue"),
            ("Burgos, Jacob M.", "Boy", "Allowed", "Blue"),
            ("Clarin, Jeremiah A.", "Boy", "Allowed", "Blue"),
            ("Dinglasan, Shant Kristoff R.", "Boy", "Allowed", "Blue"),
            ("Estenor, Arjae Luke L.", "Boy", "Allowed", "Blue"),
            ("Galleto, Neil Gabriel V.", "Boy", "Allowed", "Blue"),
            ("Garin, Yancy Ric", "Boy", "Allowed", "Blue"),
            ("Guazo, Ashton Bryce S.", "Boy", "Allowed", "Blue"),
            ("Lagmay, Alexander III.", "Boy", "Allowed", "Blue"),
            ("Manalocon, Yusof M.", "Boy", "Allowed", "Blue"),
            ("Martus, Zane Jared Rui D.", "Boy", "Allowed", "Blue"),
            ("Rosal, Alexis James", "Boy", "Allowed", "Blue"),
            ("Salangsang, Lyric Lynce V.", "Boy", "Allowed at the waiting area", "White"),
            ("Salubre, Andreau Louis I.", "Boy", "Allowed", "Blue"),
            ("Soledad, Japheth M.", "Boy", "Allowed", "Blue"),
            ("Tan, Joseph N.", "Boy", "Allowed", "Blue"),
            ("Torres, Marcus Mikael Z.", "Boy", "Allowed", "Blue"),
            ("Go, Meryl Quisa Q.", "Girl", "Allowed", "Blue"),
            ("Gonzales, Althea Grace C.", "Girl", "Allowed", "Blue"),
            ("Lagamayo, Philyn W.", "Girl", "Allowed", "Blue"),
            ("Lu, Caileen Gale K.", "Girl", "Allowed", "Blue"),
            ("Olarte, Raine Alexa K.", "Girl", "Allowed", "Blue"),
            ("Pamisa, Ma. Hynnrie Pia P.", "Girl", "Allowed", "Blue"),
            ("Taboac, Kohana M.", "Girl", "Allowed", "Blue"),
            ("Tan, Hanna Jasmine T.", "Girl", "Allowed", "Blue"),
            ("Tan, Micaella Angela S.", "Girl", "Allowed", "Blue"),
            ("Teng, Julianne Lyciel T.", "Girl", "Allowed", "Blue"),
        ]
    },
    # Other Lists / Additional Sections from scanned batches
    {
        "grade_section": "Class Records – Rezil Jaya L. Manuel",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "Rezil Jaya L. Manuel, LPT",
        "students": [
            ("Loyola, Ia D.", "Girl", "Allowed", "Blue"),
            ("Malveaux, Ariannah A.", "Girl", "Allowed", "Blue"),
            ("Sevilla, Jellana Kim S.", "Girl", "Unmarked", "-"),
            ("Solde, Aisaia Theone H.", "Girl", "Allowed", "Blue"),
            ("Tan, Bea Mei Alexiz", "Girl", "Allowed", "Blue"),
        ]
    },
    {
        "grade_section": "Class Records – Herbert T. Santos",
        "school_year": "S.Y. 2026 – 2027",
        "adviser": "Herbert T. Santos, LPT",
        "students": [
            ("Villegas, Chaira Mari C.", "Girl", "Allowed", "Blue"),
        ]
    }
]

wb = openpyxl.Workbook()

# Setup styles
header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

sub_header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
sub_header_font = Font(name="Calibri", size=11, bold=True, color="1F4E79")

title_font = Font(name="Calibri", size=14, bold=True, color="1F4E79")
subtitle_font = Font(name="Calibri", size=10, italic=True, color="595959")

thin_border = Border(
    left=Side(style='thin', color='D3D3D3'),
    right=Side(style='thin', color='D3D3D3'),
    top=Side(style='thin', color='D3D3D3'),
    bottom=Side(style='thin', color='D3D3D3')
)

header_border = Border(
    left=Side(style='thin', color='1F4E79'),
    right=Side(style='thin', color='1F4E79'),
    top=Side(style='medium', color='1F4E79'),
    bottom=Side(style='medium', color='1F4E79')
)

status_fills = {
    "Allowed": PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid"),
    "Allowed at the waiting area": PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid"),
    "Not allowed": PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid"),
    "Unmarked": PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid"),
}

status_fonts = {
    "Allowed": Font(name="Calibri", size=10, bold=True, color="375623"),
    "Allowed at the waiting area": Font(name="Calibri", size=10, bold=True, color="7F6000"),
    "Not allowed": Font(name="Calibri", size=10, bold=True, color="C65911"),
    "Unmarked": Font(name="Calibri", size=10, color="595959"),
}

# -------------------------------------------------------------
# 1. SUMMARY DASHBOARD SHEET
# -------------------------------------------------------------
ws_summary = wb.active
ws_summary.title = "Summary Dashboard"
ws_summary.views.sheetView[0].showGridLines = True

ws_summary["A1"] = "GENERAL SANTOS HOPE CHRISTIAN SCHOOL"
ws_summary["A1"].font = title_font
ws_summary["A2"] = "Lunch Break Access Summary Dashboard"
ws_summary["A2"].font = Font(name="Calibri", size=12, bold=True, color="2F5597")

headers_summary = [
    "Grade & Section", "School Year", "Class Adviser", 
    "Total Students", "Allowed (Blue)", "Allowed Waiting Area (White)", 
    "Not Allowed (Red)", "Unmarked"
]

row_idx = 4
for col_idx, h in enumerate(headers_summary, 1):
    cell = ws_summary.cell(row=row_idx, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = header_border
ws_summary.row_dimensions[row_idx].height = 28

total_all = 0
total_allowed = 0
total_waiting = 0
total_not_allowed = 0
total_unmarked = 0

for sec in data:
    row_idx += 1
    cnt_total = len(sec["students"])
    cnt_allowed = sum(1 for s in sec["students"] if s[2] == "Allowed")
    cnt_waiting = sum(1 for s in sec["students"] if s[2] == "Allowed at the waiting area")
    cnt_not_allowed = sum(1 for s in sec["students"] if s[2] == "Not allowed")
    cnt_unmarked = sum(1 for s in sec["students"] if s[2] == "Unmarked")

    total_all += cnt_total
    total_allowed += cnt_allowed
    total_waiting += cnt_waiting
    total_not_allowed += cnt_not_allowed
    total_unmarked += cnt_unmarked

    ws_summary.cell(row=row_idx, column=1, value=sec["grade_section"]).alignment = Alignment(horizontal="left", vertical="center")
    ws_summary.cell(row=row_idx, column=2, value=sec["school_year"]).alignment = Alignment(horizontal="center", vertical="center")
    ws_summary.cell(row=row_idx, column=3, value=sec["adviser"] if sec["adviser"] else "—").alignment = Alignment(horizontal="left", vertical="center")
    ws_summary.cell(row=row_idx, column=4, value=cnt_total).alignment = Alignment(horizontal="center", vertical="center")
    
    c_allow = ws_summary.cell(row=row_idx, column=5, value=cnt_allowed)
    c_allow.alignment = Alignment(horizontal="center", vertical="center")
    c_allow.fill = PatternFill(start_color="EBF1E5", fill_type="solid")
    
    c_wait = ws_summary.cell(row=row_idx, column=6, value=cnt_waiting)
    c_wait.alignment = Alignment(horizontal="center", vertical="center")
    c_wait.fill = PatternFill(start_color="FFF8E7", fill_type="solid")
    
    c_not = ws_summary.cell(row=row_idx, column=7, value=cnt_not_allowed)
    c_not.alignment = Alignment(horizontal="center", vertical="center")
    c_not.fill = PatternFill(start_color="FDF2E9", fill_type="solid")
    
    ws_summary.cell(row=row_idx, column=8, value=cnt_unmarked).alignment = Alignment(horizontal="center", vertical="center")

    for col_idx in range(1, 9):
        ws_summary.cell(row=row_idx, column=col_idx).border = thin_border
    ws_summary.row_dimensions[row_idx].height = 20

# Add Total Row
row_idx += 1
ws_summary.cell(row=row_idx, column=1, value="TOTAL").font = Font(name="Calibri", size=11, bold=True)
ws_summary.cell(row=row_idx, column=4, value=total_all).font = Font(name="Calibri", size=11, bold=True)
ws_summary.cell(row=row_idx, column=5, value=total_allowed).font = Font(name="Calibri", size=11, bold=True, color="276A3C")
ws_summary.cell(row=row_idx, column=6, value=total_waiting).font = Font(name="Calibri", size=11, bold=True, color="8D6B00")
ws_summary.cell(row=row_idx, column=7, value=total_not_allowed).font = Font(name="Calibri", size=11, bold=True, color="C00000")
ws_summary.cell(row=row_idx, column=8, value=total_unmarked).font = Font(name="Calibri", size=11, bold=True)

for col_idx in range(1, 9):
    c = ws_summary.cell(row=row_idx, column=col_idx)
    c.fill = sub_header_fill
    c.border = Border(top=Side(style="thin", color="1F4E79"), bottom=Side(style="double", color="1F4E79"))
    if col_idx >= 4:
        c.alignment = Alignment(horizontal="center", vertical="center")
ws_summary.row_dimensions[row_idx].height = 24

for col in ws_summary.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_summary.column_dimensions[col_letter].width = max(max_len + 4, 14)


# -------------------------------------------------------------
# 2. MASTER LIST SHEET
# -------------------------------------------------------------
ws_master = wb.create_sheet(title="Master Student List")
ws_master.views.sheetView[0].showGridLines = True

ws_master["A1"] = "GENERAL SANTOS HOPE CHRISTIAN SCHOOL"
ws_master["A1"].font = title_font
ws_master["A2"] = "Complete Student Lunch Break Access Master Directory"
ws_master["A2"].font = Font(name="Calibri", size=12, bold=True, color="2F5597")

headers_master = [
    "No.", "Grade & Section", "School Year", "Class Adviser", 
    "Student Name", "Gender", "Access Status", "Badge / Marker Color"
]

row_idx = 4
for col_idx, h in enumerate(headers_master, 1):
    cell = ws_master.cell(row=row_idx, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = header_border
ws_master.row_dimensions[row_idx].height = 28

global_num = 1
for sec in data:
    for s_name, gender, status, badge in sec["students"]:
        row_idx += 1
        ws_master.cell(row=row_idx, column=1, value=global_num).alignment = Alignment(horizontal="center", vertical="center")
        ws_master.cell(row=row_idx, column=2, value=sec["grade_section"]).alignment = Alignment(horizontal="left", vertical="center")
        ws_master.cell(row=row_idx, column=3, value=sec["school_year"]).alignment = Alignment(horizontal="center", vertical="center")
        ws_master.cell(row=row_idx, column=4, value=sec["adviser"] if sec["adviser"] else "—").alignment = Alignment(horizontal="left", vertical="center")
        ws_master.cell(row=row_idx, column=5, value=s_name).alignment = Alignment(horizontal="left", vertical="center")
        ws_master.cell(row=row_idx, column=6, value=gender if gender else "—").alignment = Alignment(horizontal="center", vertical="center")
        
        c_status = ws_master.cell(row=row_idx, column=7, value=status)
        c_status.alignment = Alignment(horizontal="center", vertical="center")
        c_status.fill = status_fills.get(status, PatternFill())
        c_status.font = status_fonts.get(status, Font())

        c_badge = ws_master.cell(row=row_idx, column=8, value=badge)
        c_badge.alignment = Alignment(horizontal="center", vertical="center")

        for col_idx in range(1, 9):
            ws_master.cell(row=row_idx, column=col_idx).border = thin_border
        ws_master.row_dimensions[row_idx].height = 19
        global_num += 1

for col in ws_master.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws_master.column_dimensions[col_letter].width = max(max_len + 4, 12)


# -------------------------------------------------------------
# 3. INDIVIDUAL SECTION SHEETS
# -------------------------------------------------------------
for sec in data:
    # Shorten sheet title to fit Excel limit (31 chars)
    sheet_title = sec["grade_section"].replace("Grade ", "G").replace("Class Records – ", "").replace(" – ", " ")[:31]
    ws_sec = wb.create_sheet(title=sheet_title)
    ws_sec.views.sheetView[0].showGridLines = True

    ws_sec["A1"] = f"{sec['grade_section']} — {sec['school_year']}"
    ws_sec["A1"].font = title_font
    ws_sec["A2"] = f"Class Adviser: {sec['adviser'] if sec['adviser'] else 'N/A'}"
    ws_sec["A2"].font = subtitle_font

    headers_sec = ["No.", "Student Name", "Gender", "Access Status", "Badge / Marker Color"]
    row_idx = 4
    for col_idx, h in enumerate(headers_sec, 1):
        cell = ws_sec.cell(row=row_idx, column=col_idx, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = header_border
    ws_sec.row_dimensions[row_idx].height = 26

    sec_num = 1
    for s_name, gender, status, badge in sec["students"]:
        row_idx += 1
        ws_sec.cell(row=row_idx, column=1, value=sec_num).alignment = Alignment(horizontal="center", vertical="center")
        ws_sec.cell(row=row_idx, column=2, value=s_name).alignment = Alignment(horizontal="left", vertical="center")
        ws_sec.cell(row=row_idx, column=3, value=gender if gender else "—").alignment = Alignment(horizontal="center", vertical="center")
        
        c_status = ws_sec.cell(row=row_idx, column=4, value=status)
        c_status.alignment = Alignment(horizontal="center", vertical="center")
        c_status.fill = status_fills.get(status, PatternFill())
        c_status.font = status_fonts.get(status, Font())

        c_badge = ws_sec.cell(row=row_idx, column=5, value=badge)
        c_badge.alignment = Alignment(horizontal="center", vertical="center")

        for col_idx in range(1, 6):
            ws_sec.cell(row=row_idx, column=col_idx).border = thin_border
        ws_sec.row_dimensions[row_idx].height = 19
        sec_num += 1

    # Section Summary Footers
    row_idx += 2
    ws_sec.cell(row=row_idx, column=1, value="Section Breakdown:").font = Font(name="Calibri", size=10, bold=True)
    row_idx += 1
    ws_sec.cell(row=row_idx, column=2, value="Allowed (Blue):")
    ws_sec.cell(row=row_idx, column=3, value=sum(1 for s in sec["students"] if s[2] == "Allowed")).font = Font(bold=True)
    row_idx += 1
    ws_sec.cell(row=row_idx, column=2, value="Allowed at Waiting Area (White):")
    ws_sec.cell(row=row_idx, column=3, value=sum(1 for s in sec["students"] if s[2] == "Allowed at the waiting area")).font = Font(bold=True)
    row_idx += 1
    ws_sec.cell(row=row_idx, column=2, value="Not Allowed (Red):")
    ws_sec.cell(row=row_idx, column=3, value=sum(1 for s in sec["students"] if s[2] == "Not allowed")).font = Font(bold=True)

    for col in ws_sec.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_sec.column_dimensions[col_letter].width = max(max_len + 4, 12)

# Save to destination paths
out_path_docs = r"c:\Users\HowardKevinVelos\Documents\HOPE ID SCANNER\Student_Lunch_Break_Access_List.xlsx"
out_path_pics = r"C:\Users\HowardKevinVelos\Pictures\student access\Student_Lunch_Break_Access_List.xlsx"

wb.save(out_path_docs)
print(f"Successfully generated: {out_path_docs}")

try:
    wb.save(out_path_pics)
    print(f"Successfully saved copy in Pictures: {out_path_pics}")
except Exception as e:
    print(f"Could not save in Pictures: {e}")
